
%%writefile app.py

import io
import os
from typing import List

import streamlit as st
from google import genai
from pydantic import BaseModel, Field, field_validator
from pypdf import PdfReader
from docx import Document


# ============================================================
# 1. PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="ResumeIQ | ATS Resume Analyzer",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="expanded",
)


# ============================================================
# 2. CUSTOM CSS STYLING
# ============================================================

st.markdown("""
<style>
@import url(
'https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;600;700&family=Manrope:wght@400;500;600;700;800&display=swap'
);

html, body, [class*="css"] {
    font-family: 'DM Sans', sans-serif;
}

.stApp {
    background-color: #f6f7fb;
}

.block-container {
    padding-top: 2rem;
    padding-bottom: 3rem;
    max-width: 1180px;
}

.hero {
    padding: 2rem;
    border-radius: 24px;
    background: linear-gradient(
        120deg, #111827, #312e81, #6d28d9
    );
    color: white;
    margin-bottom: 1.5rem;
}

.hero h1 {
    color: white;
    font-family: 'Manrope', sans-serif;
    font-size: 2.4rem;
    font-weight: 800;
}

.hero p {
    color: #e0e7ff;
    font-size: 1.05rem;
}

div[data-testid="stMetric"] {
    background: white;
    border: 1px solid #e7e9f2;
    padding: 1rem;
    border-radius: 14px;
}

.stButton button {
    border-radius: 12px;
    font-weight: 700;
    min-height: 3rem;
}

section[data-testid="stSidebar"] {
    background: white;
    border-right: 1px solid #e7e9f2;
}
</style>
""", unsafe_allow_html=True)


# ============================================================
# 3. STRUCTURED RESPONSE MODEL
# ============================================================

class ResumeAnalysis(BaseModel):
    """Defines the expected Gemini response structure."""

    ats_score: int = Field(
        description="Resume-to-job match score from 0 to 100"
    )

    summary: str = Field(
        description="Concise summary of the resume's alignment"
    )

    strengths: List[str] = Field(
        description="Relevant skills and strengths supported by the resume"
    )

    weaknesses: List[str] = Field(
        description="Missing or weakly demonstrated job requirements"
    )

    recommendations: List[str] = Field(
        description="Specific and actionable resume improvements"
    )

    @field_validator("ats_score")
    @classmethod
    def validate_score(cls, score):
        """Keep the displayed score within the expected range."""
        return max(0, min(100, score))


# ============================================================
# 4. RESUME TEXT EXTRACTION
# ============================================================

def read_resume(uploaded_file) -> str:
    """
    Extract readable text from PDF, DOCX, or TXT files.
    """

    if uploaded_file is None:
        raise ValueError("Please upload a resume.")

    filename = uploaded_file.name.lower()
    file_bytes = uploaded_file.getvalue()

    if filename.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(file_bytes))

        resume_text = "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    elif filename.endswith(".docx"):
        document = Document(io.BytesIO(file_bytes))

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
        ]

        # Extract text from tables as well.
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(
                    " | ".join(cell.text for cell in row.cells)
                )

        resume_text = "\n".join(paragraphs)

    elif filename.endswith(".txt"):
        resume_text = file_bytes.decode(
            "utf-8", errors="replace"
        )

    else:
        raise ValueError(
            "Unsupported format. Upload PDF, DOCX, or TXT."
        )

    if not resume_text.strip():
        raise ValueError(
            "No readable text found. Scanned PDFs may require OCR."
        )

    return resume_text.strip()


# ============================================================
# 5. GEMINI API CONFIGURATION
# ============================================================

def get_api_key() -> str:
    """
    Retrieve the key from Streamlit Secrets or an environment
    variable. Never hard-code or display the actual API key.
    """

    try:
        secret_key = st.secrets.get("GEMINI_API_KEY", "")
    except Exception:
        secret_key = ""

    return secret_key or os.getenv("GEMINI_API_KEY", "")


# ============================================================
# 6. AI-POWERED RESUME ANALYSIS
# ============================================================

def analyze_resume(
    resume_text: str,
    job_description: str,
    model_name: str = "gemini-2.5-flash",
) -> ResumeAnalysis:
    """
    Compare resume evidence against a job description using
    Gemini and validate the response with Pydantic.
    """

    api_key = get_api_key()

    if not api_key:
        raise RuntimeError(
            "Gemini API key is missing. Configure GEMINI_API_KEY."
        )

    client = genai.Client(api_key=api_key)

    prompt = f"""
You are an ATS-style resume analysis assistant.

Analyze the resume against the specific job description.

JOB DESCRIPTION:
{job_description}

RESUME:
{resume_text}

Return the following fields:
1. ats_score: Integer from 0 to 100.
2. summary: Concise alignment summary.
3. strengths: List of matching skills and evidence.
4. weaknesses: List of missing or weakly demonstrated requirements.
5. recommendations: Specific action points for improvement.

IMPORTANT RULES:
- Treat the score as an approximate relevance estimate.
- Do not claim it is an actual employer ATS score.
- Do not invent qualifications, experience, projects, or skills.
- Distinguish missing skills from skills that are not clearly evidenced.
- Consider the requirements of entry-level candidates fairly.
- Recommend adding only truthful information.
- Return concise, practical feedback.
"""

    response = client.models.generate_content(
        model=model_name,
        contents=prompt,
        config={
            "response_mime_type": "application/json",
            "response_schema": ResumeAnalysis,
        },
    )

    # Prefer the SDK's parsed structured response when available.
    if getattr(response, "parsed", None) is not None:
        return ResumeAnalysis.model_validate(response.parsed)

    # Otherwise, validate the JSON response text.
    if not response.text:
        raise RuntimeError(
            "Gemini returned an empty response. Please retry."
        )

    return ResumeAnalysis.model_validate_json(response.text)


# ============================================================
# 7. APPLICATION SIDEBAR
# ============================================================

with st.sidebar:
    st.markdown("## 📄 ResumeIQ")
    st.caption("AI-powered resume and job matching")

    st.divider()

    model_name = st.selectbox(
        "Gemini model",
        ["gemini-2.5-flash", "gemini-2.5-pro"],
        index=0,
        help="Availability depends on your Google AI API access.",
    )

    st.markdown("### How it works")
    st.markdown("""
    1. Upload your resume.
    2. Paste the target job description.
    3. Generate the match analysis.
    4. Review gaps and recommendations.
    """)

    st.divider()

    st.caption(
        "Resume text is sent to Gemini for analysis. "
        "Upload documents only when you are comfortable "
        "with that processing."
    )


# ============================================================
# 8. MAIN HEADER
# ============================================================

st.markdown("""
<div class="hero">
    <h1>Make your next application count.</h1>
    <p>
        Compare your resume with a job description and get
        focused, actionable feedback in seconds.
    </p>
</div>
""", unsafe_allow_html=True)

col1, col2, col3 = st.columns(3)

with col1:
    st.markdown("### 🎯 Job-specific matching")
    st.caption("Evaluate your resume against a target role.")

with col2:
    st.markdown("### 🔍 Skills gap analysis")
    st.caption("Identify missing or unclear requirements.")

with col3:
    st.markdown("### 🚀 Actionable feedback")
    st.caption("Get practical, honest improvement suggestions.")

st.divider()


# ============================================================
# 9. RESUME AND JOB DESCRIPTION INPUTS
# ============================================================

left, right = st.columns([0.9, 1.1], gap="large")

with left:
    st.subheader("01 · Your resume")

    uploaded_file = st.file_uploader(
        "Upload your resume",
        type=["pdf", "docx", "txt"],
        help="Use a text-based PDF, DOCX, or TXT file.",
    )

    if uploaded_file:
        st.success(f"Ready: {uploaded_file.name}")
        st.caption(
            f"File size: {uploaded_file.size / 1024:.1f} KB"
        )

with right:
    st.subheader("02 · Target job")

    job_description = st.text_area(
        "Job description",
        height=230,
        placeholder=(
            "Paste the job title, required skills, "
            "responsibilities, and qualifications..."
        ),
    )

analyze_clicked = st.button(
    "✨ Analyze My Resume",
    type="primary",
    use_container_width=True,
)


# ============================================================
# 10. RUN THE ANALYSIS
# ============================================================

if analyze_clicked:

    if uploaded_file is None:
        st.error("Please upload a resume.")

    elif not job_description.strip():
        st.error("Please enter a job description.")

    elif len(job_description.strip()) < 40:
        st.error(
            "Please provide a more detailed job description."
        )

    else:
        try:
            with st.spinner(
                "Reading your resume and analyzing the job match..."
            ):
                resume_text = read_resume(uploaded_file)

                # Limit prompt size for practical API usage.
                resume_text = resume_text[:45000]

                result = analyze_resume(
                    resume_text,
                    job_description.strip(),
                    model_name,
                )

            st.session_state["analysis_result"] = result
            st.session_state["analysis_filename"] = uploaded_file.name

        except Exception as error:
            st.error(f"Analysis failed: {error}")
            st.info(
                "Check your API key, Gemini model access, "
                "internet connection, and resume format."
            )


# ============================================================
# 11. DISPLAY ANALYSIS RESULTS
# ============================================================

result = st.session_state.get("analysis_result")

if result:

    st.divider()
    st.header("📊 Your Resume Analysis")

    st.caption(
        f"Resume analyzed: "
        f"{st.session_state.get('analysis_filename', 'Resume')}"
    )

    score = max(0, min(100, result.ats_score))

    score_col, summary_col = st.columns([0.75, 1.65])

    with score_col:
        st.metric("Estimated ATS Match", f"{score}/100")
        st.progress(score / 100)

        if score >= 75:
            st.success("Strong alignment")
        elif score >= 50:
            st.warning("Partial alignment")
        else:
            st.info("Needs more tailoring")

    with summary_col:
        st.subheader("📝 Executive summary")
        st.write(result.summary)

        if result.recommendations:
            st.subheader("💡 Priority improvement")
            st.write(result.recommendations[0])

    st.divider()

    strengths_col, weaknesses_col = st.columns(2)

    with strengths_col:
        st.subheader("✅ Strengths")

        if result.strengths:
            for item in result.strengths:
                st.markdown(f"- {item}")
        else:
            st.write("No specific strengths returned.")

    with weaknesses_col:
        st.subheader("🔍 Gaps to address")

        if result.weaknesses:
            for item in result.weaknesses:
                st.markdown(f"- {item}")
        else:
            st.write("No specific gaps returned.")

    st.divider()
    st.subheader("🚀 Recommendations")

    for index, item in enumerate(
        result.recommendations, start=1
    ):
        st.markdown(f"{index}. {item}")

    # Let users download the structured analysis.
    st.download_button(
        label="⬇️ Download Analysis as JSON",
        data=result.model_dump_json(indent=2),
        file_name="resume_analysis.json",
        mime="application/json",
    )

    st.caption(
        "The match score is an AI-generated estimate, not a "
        "guarantee of employer ATS results. Verify every "
        "recommendation before changing your resume."
    )

else:
    st.info(
        "Upload your resume and paste a job description to "
        "generate your personalized analysis."
    )
